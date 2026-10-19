#!/usr/bin/env python3
"""Identity gate for the concentration policy (golden mining r2, sheet
24_ConcentrazioneMensile). Mirrors :mod:`test_fixed_weekday_identity` minus the
draft tier (the proven artifact is the CORRECTED table in
``claudedocs/spalmone-2026-27/golden-mining-r2-2026-08-16.md``; the first pass
of that doc counted stage-day hours as aula+lab capacity - the correction note
is in the doc itself). Two tiers:

1. **Sanity invariants:** the derivation returns exactly the 6 STRONG rows the
   corrected mining table measured; representative rows match byte for byte
   (G02 Barberia febbraio 18h, G12 Servizio tecnico aprile 32h); every row
   respects both floors (>= 4x expected AND >= 25% of the year in the month);
   and the two Cad G08 artifacts of the buggy baseline are ABSENT (they sit
   at 3.1x/2.8x under the corrected one).
2. **Round-trip emission:** build the workbook, reload it, read
   24_ConcentrazioneMensile BY HEADER NAME and check its rows match the
   derivation (catches column misalignment or dropped rows in the emitter).

Drift handling: if the golden, the calendar sheets or the thresholds change,
the live derivation changes and this gate fails LOUD (exit != 0). That is the
signal to re-measure and update the mining doc first. Run with the venv python::

    scripts/workbook/.venv/bin/python scripts/workbook/test_concentration_identity.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import concentration_policy as CP
from generate_workbook import build_workbook

EXPECTED_ROWS = 6  # corrected mining table (golden-mining-r2-2026-08-16.md, section B)

# The corrected doc's named rows, pinned: (classe, materia, mese, oreNelMese, oreAnno).
PINNED_ROWS = [
    ("G02", "Servizio stilistico", "gennaio", 18, 40),
    ("G02", "Barberia", "febbraio", 18, 36),
    ("G12", "Servizio tecnico e stilistico", "aprile", 32, 46),
    ("G04", "Recuperi ed Approfondimenti", "aprile", 28, 50),
]

# Artifacts of the FIRST (buggy) baseline: S-day hours counted as aula+lab.
MUST_BE_ABSENT = [
    ("G08", "Cad", "febbraio"),   # 3.1x under the corrected baseline
    ("G08", "Cad", "maggio"),     # 2.8x
]


def fail(msg: str) -> None:
    print(f"FAIL {msg}")
    sys.exit(1)


def main() -> None:
    rows = CP.concentration_pairs()

    # ── tier 1: sanity invariants ────────────────────────────────────────
    if len(rows) != EXPECTED_ROWS:
        fail(f"expected {EXPECTED_ROWS} strong rows, got {len(rows)}: "
             f"{[(r['classe'], r['materia'], r['mese']) for r in rows]}")
    for r in rows:
        if r["rapporto"] < CP.RATIO_FLOOR:
            fail(f"{r['classe']}/{r['materia']}/{r['mese']} ratio {r['rapporto']:.2f} < "
                 f"{CP.RATIO_FLOOR}: the floor is part of the declared rule")
        if r["quotaAnno"] < CP.SHARE_FLOOR:
            fail(f"{r['classe']}/{r['materia']}/{r['mese']} share {r['quotaAnno']:.2f} < "
                 f"{CP.SHARE_FLOOR}: the floor is part of the declared rule")
    present = {(r["classe"], r["materia"], r["mese"]) for r in rows}
    for pin in PINNED_ROWS:
        match = [r for r in rows if (r["classe"], r["materia"], r["mese"]) == pin[:3]]
        if not match:
            fail(f"pinned row missing: {pin}")
        if match[0]["oreNelMese"] != pin[3] or match[0]["oreAnno"] != pin[4]:
            fail(f"pinned row drifted: {pin} -> {match[0]}")
    for absent in MUST_BE_ABSENT:
        if absent in present:
            fail(f"buggy-baseline artifact present: {absent} (S-day hours must not count"
                 " as aula+lab capacity)")
    print(f"tier 1 OK: {EXPECTED_ROWS} strong rows, floors hold, pins match,"
          f" {len(MUST_BE_ABSENT)} artifacts absent")

    # ── tier 2: round-trip emission ──────────────────────────────────────
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "wb.xlsx"
        wb = build_workbook()
        wb.save(path)
        reloaded = load_workbook(path, read_only=True)
        if "24_ConcentrazioneMensile" not in reloaded.sheetnames:
            fail("sheet 24_ConcentrazioneMensile missing from the workbook")
        ws = reloaded["24_ConcentrazioneMensile"]
        it = ws.iter_rows(values_only=True)
        header = list(next(it))
        col = {name: header.index(name) for name in
               ("classe", "materia", "mese", "oreNelMese", "rapporto", "oreAnno")}
        sheet_rows = []
        for values in it:
            if values[col["classe"]] is None:
                continue
            sheet_rows.append((values[col["classe"]], values[col["materia"]],
                               values[col["mese"]], values[col["oreNelMese"]],
                               values[col["oreAnno"]]))
        reloaded.close()
    expected_rows = sorted((r["classe"], r["materia"], r["mese"],
                            r["oreNelMese"], r["oreAnno"]) for r in rows)
    if sorted(sheet_rows) != expected_rows:
        fail(f"round-trip drift: sheet has {len(sheet_rows)} rows,"
             f" derivation has {len(expected_rows)}")
    print(f"tier 2 OK: round-trip emission matches ({len(sheet_rows)} rows)")
    print("PASS concentration identity")


if __name__ == "__main__":
    main()
