#!/usr/bin/env python3
"""Byte-identity gate for the theory-block policy (TASK-471, sheet 22_TeoriaBlocco).

The contract: the policy the workbook derives LIVE from the golden must reproduce
the Fedora-verified draft (``task471-theory-policy-draft.csv``), row for row. The
draft is the proven artifact (196 (classe, materia) pairs, split 157/31/8, with
the "mecaniche" typo bug already corrected by importing the authoritative
``is_theory``). This gate proves the live derivation in
:mod:`theory_block_policy` stays faithful to it, so a regenerated workbook always
carries the same declared input referente is asked to confirm.

Why SET-based, not line-based: the draft's row ORDER is presentation (it groups by
classification within each class); the load-bearing content is the SET of
(classe, materia, evidence, policy) rows. Comparing as sets keeps the gate honest
about what is proven (the values) and lets the sheet emitter choose its own
deterministic sort without a fragile line-order coupling to the draft.

Three tiers, kept apart so the reader knows what is proven:

1. **Derivation == draft (set-based).** :func:`theory_block_policy.theory_block_policy`
   live vs the committed CSV. THIS is the load-bearing byte-identity link.
2. **Sanity invariants (independent of the CSV):** 196 pairs, 6800 theory hours,
   the 157/31/8 split. They catch a logic regression even if both sides drifted
   together.
3. **Round-trip emission:** build the workbook, reload it, read 22_TeoriaBlocco BY
   HEADER NAME and check its rows match the derivation. Tolerates a benign column
   reorder, but catches a value/column misalignment or a dropped row in
   ``sheet_22_teoria_blocco`` that the derivation-only check cannot see. Mirrors
   :mod:`test_materie_byte_identity`.

Drift handling: if the golden ``seed/lessons.csv`` changes, the live policy
changes and this gate fails LOUD (exit != 0 with the divergent rows). That is the
signal to regenerate and re-commit the draft CSV (the new proven artifact) and
re-run. Silent absorption is exactly what the gate exists to prevent.

Run with the venv python (the round-trip tier needs openpyxl + build_workbook)::

    scripts/workbook/.venv/bin/python scripts/workbook/test_theory_block_identity.py
"""

from __future__ import annotations

import csv
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

import theory_block_policy as TBP
from generate_workbook import build_workbook

DRAFT_CSV = (Path(__file__).resolve().parents[2] / "claudedocs" / "spalmone-2026-27"
             / "task471-theory-policy-draft.csv")


def _row_key(classe, materia, hours, active, days2, avg, policy):
    """The comparable identity of a row. avg is canonicalized to the same 2-decimal
    string the draft carries, so 1.1538 (live) and "1.15" (draft) compare equal."""
    return (classe, materia, int(hours), int(active), int(days2), f"{float(avg):.2f}", policy)


def _load_draft() -> set:
    """Parse the proven draft CSV into a set of row keys."""
    out = set()
    with DRAFT_CSV.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            out.add(_row_key(r["group"], r["subject"], r["theory_hours"], r["active_days"],
                             r["days_2plus"], r["avg_hours_per_active_day"], r["block_policy"]))
    return out


def _sheet_row_set() -> set:
    """Build the workbook, reload it, read 22_TeoriaBlocco BY HEADER NAME into a set
    of (classe, materia, oreTeoria, giorniAttivi, giorniConBlocco, bloccoTeoria).

    Reading by header (not position) tolerates a benign reorder but still catches a
    value/column misalignment on the six identity-bearing columns it reads: if the
    emitter swaps two appends, the wrong value lands under a header and the set
    diverges from the derivation. The three display columns (mediaOreGiorno,
    origine, note) are intentionally not checked here; the classification they
    rest on is already pinned by tier 1."""
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "gate.xlsx"
        build_workbook().save(path)
        ws = load_workbook(path)["22_TeoriaBlocco"]
    names = ["classe", "materia", "oreTeoria", "giorniAttivi", "giorniConBlocco", "bloccoTeoria"]
    col = {str(ws.cell(row=1, column=c).value): c for c in range(1, ws.max_column + 1)}
    out = set()
    for r in range(2, ws.max_row + 1):
        out.add(tuple(ws.cell(row=r, column=col[name]).value for name in names))
    return out


def main() -> int:
    errors = []

    # Derive ONCE from the golden; the live set, the sanity invariants and the
    # round-trip expectation all read from this same list (no second golden parse).
    rows = TBP.theory_block_policy()
    live = {_row_key(r["classe"], r["materia"], r["oreTeoria"], r["giorniAttivi"],
                     r["giorniConBlocco"], r["avg"], r["bloccoTeoria"]) for r in rows}
    draft = _load_draft()

    # 1. Byte-identity (set-based): live derivation == proven draft.
    if live != draft:
        missing = sorted(draft - live)   # in draft, not reproduced live
        extra = sorted(live - draft)     # produced live, not in draft
        for k in missing[:20]:
            errors.append(f"MANCA (nel draft, non riprodotto live): {k}")
        for k in extra[:20]:
            errors.append(f"EXTRA (prodotto live, non nel draft): {k}")
        if len(missing) > 20 or len(extra) > 20:
            errors.append(f"...({len(missing)} mancanti, {len(extra)} extra in totale)")

    # 2. Sanity invariants (the TASK-471 draft headline, independent of the CSV).
    if len(rows) != 196:
        errors.append(f"pairs = {len(rows)}, attese 196 (golden o is_theory cambiati)")
    total_hours = sum(r["oreTeoria"] for r in rows)
    if total_hours != 6800:
        errors.append(f"theory hours = {total_hours}, attesi 6800")
    split = Counter(r["bloccoTeoria"] for r in rows)
    for policy, want in {TBP.STRETTA: 157, TBP.MISTA: 31, TBP.BLOCCO: 8}.items():
        got = split.get(policy, 0)
        if got != want:
            errors.append(f"split {policy!r} = {got}, atteso {want}")

    # 3. Round-trip: the SHEET the school opens carries these same rows.
    sheet = _sheet_row_set()
    expected = {(r["classe"], r["materia"], r["oreTeoria"], r["giorniAttivi"],
                 r["giorniConBlocco"], r["bloccoTeoria"]) for r in rows}
    if sheet != expected:
        errors.append("foglio 22_TeoriaBlocco (round-trip) != derivazione "
                      f"(missing={sorted(expected - sheet)[:5]}, "
                      f"extra={sorted(sheet - expected)[:5]})")

    if errors:
        print(f"FAIL: gate theory-block identity NON superato ({len(errors)} problemi)")
        for e in errors:
            print("  - " + e)
        return 1

    print(f"OK: theory-block policy live (theory_block_policy) == draft Fedora-verificato "
          f"({len(live)} coppie (classe,materia), set-based); sanity 196/6800h/split "
          f"157-31-8; round-trip foglio 22_TeoriaBlocco OK. Input dichiarato fedele.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
