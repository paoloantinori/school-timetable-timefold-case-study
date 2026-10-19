#!/usr/bin/env python3
"""Identity gate for the fixed-weekday policy (golden mining, sheet 23_GiornoFisso).

Mirrors :mod:`test_theory_block_identity` minus the draft-CSV tier (there is no
Fedora-verified draft for this sheet: the proven artifact is the mining analysis in
``claudedocs/spalmone-2026-27/golden-mining-2026-08-15.md``). Two tiers:

1. **Sanity invariants (independent of everything else):** the derivation returns the
   40 GENUINELY NEW pairs the mining doc measured (67 fixed total, 27 implied by the
   cadence reference and excluded); representative rows match the doc's findings byte
   for byte (G03 Disegno on Martedi, G02 Servizio tecnico e stilistico on
   Giovedi); and NO returned pair is cadence-implied (the deducibility filter holds).
2. **Round-trip emission:** build the workbook, reload it, read 23_GiornoFisso BY
   HEADER NAME and check its rows match the derivation (catches value/column
   misalignment or dropped rows in the sheet emitter).

Drift handling: if the golden ``seed/lessons.csv`` or the cadence reference changes,
the live derivation changes and this gate fails LOUD (exit != 0). That is the signal
to re-measure, update the mining doc (the proven artifact) and the expected numbers
here. Silent absorption is what the gate exists to prevent.

Run with the venv python (the round-trip tier needs openpyxl)::

    scripts/workbook/.venv/bin/python scripts/workbook/test_fixed_weekday_identity.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import fixed_weekday_policy as FWP
from generate_workbook import build_workbook

EXPECTED_PAIRS = 40   # golden-mining-2026-08-15.md: 67 fixed, 27 cadence-implied, 40 new

# The mining doc's named findings, pinned: (classe, materia, giorno, ore).
PINNED_ROWS = [
    ("G03", "Disegno", "Martedì", 55),
    ("G03", "Elettrotecnica", "Mercoledì", 40),
    ("G02", "Servizio tecnico e stilistico", "Giovedì", 60),
    ("G11", "Storia", "Venerdì", 20),
    ("G11", "Normativa", "Martedì", 20),
]

failures: list[str] = []


def check(cond: bool, msg: str) -> None:
    if not cond:
        failures.append(msg)


def main() -> int:
    # ── tier 1: sanity invariants on the derivation ────────────────────────
    rows = FWP.fixed_weekday_pairs()
    check(len(rows) == EXPECTED_PAIRS,
          f"pair count {len(rows)} != {EXPECTED_PAIRS} (golden changed? re-measure)")

    by_key = {(r["classe"], r["materia"]): r for r in rows}
    for classe, materia, giorno, ore in PINNED_ROWS:
        r = by_key.get((classe, materia))
        check(r is not None, f"pinned pair missing: {classe} / {materia}")
        if r:
            check(r["giorno"] == giorno and r["ore"] == ore,
                  f"pinned pair drifted: {classe} / {materia} = {r['giorno']} {r['ore']}h "
                  f"(expected {giorno} {ore}h)")

    cadence = FWP._cadence_reference()
    leaked = [r for r in rows if FWP._implied_by_cadence(cadence, r["classe"], r["materia"])]
    check(not leaked,
          f"deducibility filter broken: {len(leaked)} cadence-implied pairs leaked into the sheet")

    # ── tier 2: round-trip emission ────────────────────────────────────────
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "wb.xlsx"
        wb = build_workbook()
        wb.save(path)
        reloaded = load_workbook(path, data_only=True)
        check("23_GiornoFisso" in reloaded.sheetnames, "sheet 23_GiornoFisso missing from workbook")
        if "23_GiornoFisso" in reloaded.sheetnames:
            ws = reloaded["23_GiornoFisso"]
            header = [c.value for c in ws[1]]
            col = {name: i for i, name in enumerate(header)}
            emitted = set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[col["classe"]] is None:
                    continue
                emitted.add((row[col["classe"]], row[col["materia"]],
                             row[col["giorno"]], row[col["ore"]]))
            derived = {(r["classe"], r["materia"], r["giorno"], r["ore"]) for r in rows}
            check(emitted == derived,
                  f"round-trip mismatch: sheet {len(emitted)} rows vs derivation {len(derived)}; "
                  f"only-sheet={sorted(emitted - derived)[:3]}; only-derivation={sorted(derived - emitted)[:3]}")

    if failures:
        print("FAIL: fixed-weekday identity gate")
        for f in failures:
            print("  -", f)
        return 1
    print(f"OK: fixed-weekday identity gate ({EXPECTED_PAIRS} pairs, "
          f"{len(PINNED_ROWS)} pinned rows, deducibility filter clean, round-trip exact)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
