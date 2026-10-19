#!/usr/bin/env python3
"""Byte-identity gate for 08_Materie (TASK-445 Tappa 1, proposal 5.2 + audit C1).

The contract: the flags the workbook proposes for 08_Materie must match what the
system computes today, on the 266 real incarichi rows. Three tiers of check,
kept honestly apart so the reader knows what is actually proven:

1. **Taxonomy flags (laboratorio / professionalizzante / tagAula) - STRONG, live.**
   These come from ``seed/subjects.csv`` (TASK-437 SubjectRegistry), itself the
   byte-identical extraction of the old DataGenerator prefix matchers and guarded
   against the Java side by ``SubjectRegistryTest``. This gate cross-checks the
   workbook back against the live CSV, so a regenerated workbook always matches
   the system's extracted truth. THIS is the load-bearing byte-identity link.

2. **Timing columns (capOreGiorno / settimanaInizioMax / settimanaInizioMin /
   primaDelloStage) - WEAK, self-consistency only.** These come from solver-side
   constants still in Java; nothing in this pipeline reads Java, so the gate
   compares the workbook against the documented snapshots in :mod:`sources` -
   the SAME literals the generator writes from. This tier can only catch a
   generator-logic bug or an openpyxl round-trip data-loss bug, NOT a stale
   transcription if the Java constants later change. Making this link mechanical
   (a Java-side test like ``SubjectRegistryTest``, or codegen from the Java
   source) is deliberate future work; until then the ``# SRC:`` comments in
   sources.py are documentation, not enforcement.

3. **Coverage + cardinality invariants (independent).** Every teaching subject in
   incarichi.csv (real teachers only) is classified in 08_Materie; the workbook
   has exactly the subjects.csv rows (no more, no less); and the seed-shape
   counts (266 real incarichi, 52 subjects) still hold.

Failure is VISIBLE: every divergent cell prints ``subject | field | expected |
actual`` and the process exits non-zero. Run with the venv python::

    scripts/workbook/.venv/bin/python scripts/workbook/test_materie_byte_identity.py
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import sources as S
from generate_workbook import build_workbook

# 08_Materie column index (1-based), matching the header order in generate_workbook.
COL = {
    "materia": 1, "laboratorio": 2, "professionalizzante": 3, "tagAula": 4,
    "capOreGiorno": 6, "settimanaInizioMax": 7, "settimanaInizioMin": 8,
    "primaDelloStage": 9,
}


def _norm_bool(cell) -> str:
    """Canonical SI/NO from a workbook cell (tolerant of bool/str/None)."""
    if cell is None or cell == "":
        return ""
    if isinstance(cell, bool):
        return "SI" if cell else "NO"
    s = str(cell).strip().upper()
    if s in ("SI", "TRUE", "S"):
        return "SI"
    if s in ("NO", "FALSE", "N"):
        return "NO"
    return s


def _norm_tags(cell) -> str:
    """Canonical ';'-joined sorted lowercase tag list (order-independent compare."""
    if cell is None or cell == "":
        return ""
    parts = [t.strip().lower() for t in str(cell).split(";") if t.strip()]
    return ";".join(sorted(parts))


def _norm_scalar(cell):
    """Empty stays empty; ints/strings stringified and stripped."""
    if cell is None:
        return ""
    return str(cell).strip()


def expected_for_subject(norm_subject: str, csv_row: dict) -> dict:
    """What 08_Materie SHOULD contain for this subject, from the live sources."""
    return {
        "laboratorio": "SI" if csv_row["roomTags"] else "NO",
        "professionalizzante": "SI" if csv_row["professionalizzanti"] else "NO",
        "tagAula": ";".join(sorted(t.lower() for t in csv_row["roomTags"])),
        "capOreGiorno": "0" if norm_subject in S.M21_EXEMPT_SUBJECTS else "",
        "settimanaInizioMax": str(S.T20_EARLY_MAX_WEEK) if norm_subject in S.T20_EARLY_SUBJECTS else "",
        "settimanaInizioMin": str(S.T20_LATE_MIN_WEEK) if norm_subject in S.T20_LATE_SUBJECTS else "",
        "primaDelloStage": "SI" if norm_subject in S.PRIMA_DELLO_STAGE_SUBJECTS else "NO",
    }


def actual_from_workbook(ws) -> dict[str, dict]:
    """Read 08_Materie body rows -> {normalized_subject: {field: raw_cell}}."""
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=False):
        materia = row[COL["materia"] - 1].value
        if materia is None or str(materia).strip() == "":
            continue
        norm = S.normalize_subject(str(materia).strip())
        out[norm] = {f: row[COL[f] - 1].value for f in COL if f != "materia"}
    return out


def main() -> int:
    # Build the workbook and round-trip it through save+reload, so the gate
    # proves the FILE (what the school opens), not just the in-memory object.
    wb = build_workbook()
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "gate.xlsx"
        wb.save(path)
        wb = load_workbook(path)
    ws = wb["08_Materie"]
    actual = actual_from_workbook(ws)

    expected_csv = {S.normalize_subject(s["subject"]): s for s in S.load_subjects()}
    errors = []

    # 1. Cardinality: workbook 08_Materie == subjects.csv, exact set.
    if set(actual) != set(expected_csv):
        missing = sorted(set(expected_csv) - set(actual))
        extra = sorted(set(actual) - set(expected_csv))
        if missing:
            errors.append(f"08_Materie MANCA materie presenti in subjects.csv: {missing}")
        if extra:
            errors.append(f"08_Materie ha materie ASSENTE in subjects.csv (extra): {extra}")

    # 2. Field-by-field byte identity for every shared subject.
    for norm, csv_row in expected_csv.items():
        if norm not in actual:
            continue  # already reported as missing above
        exp = expected_for_subject(norm, csv_row)
        act_cells = actual[norm]
        field_canon = {
            "laboratorio": _norm_bool(act_cells["laboratorio"]),
            "professionalizzante": _norm_bool(act_cells["professionalizzante"]),
            "tagAula": _norm_tags(act_cells["tagAula"]),
            "capOreGiorno": _norm_scalar(act_cells["capOreGiorno"]),
            "settimanaInizioMax": _norm_scalar(act_cells["settimanaInizioMax"]),
            "settimanaInizioMin": _norm_scalar(act_cells["settimanaInizioMin"]),
            "primaDelloStage": _norm_bool(act_cells["primaDelloStage"]),
        }
        for field, want in exp.items():
            got = field_canon[field]
            if got != want:
                errors.append(f"DIVERGENZA | materia={norm!r} campo={field} "
                              f"atteso={want!r} trovato={got!r}")

    # 3. Coverage invariant (audit C1): every teaching subject in incarichi.csv
    #    (real teachers only) is classified in 08_Materie.
    real_subjects = set()
    for row in S.load_incarichi():
        if row["teacher"] in S.PSEUDO_COURSE_TEACHERS:
            continue
        # subject "None" is the literal on the Totale-ore-corso summary rows
        # (already excluded by the teacher filter above); kept as a belt-and-
        # braces guard against any future real-teacher row with an empty subject.
        if row["subject"] and row["subject"] != "None":
            real_subjects.add(S.normalize_subject(row["subject"]))
    unclassified = sorted(real_subjects - set(actual))
    if unclassified:
        errors.append(f"Materie di insegnamento reali NON classificate in 08_Materie: {unclassified}")

    # 4. Sanity: the count assertions the TASK-437 guard relies on still hold
    #    (the proposal's "266 real incarichi" and "52 subjects").
    real_incarichi = sum(1 for r in S.load_incarichi()
                         if r["teacher"] not in S.PSEUDO_COURSE_TEACHERS and r["subject"] and r["subject"] != "None")
    if real_incarichi != 266:
        errors.append(f"incarichi reali = {real_incarichi}, attesi 266 (pseudo-teacher set o seed cambiato)")
    if len(expected_csv) != 52:
        errors.append(f"subjects.csv = {len(expected_csv)} righe, attese 52")

    if errors:
        print("FAIL: gate di byte-identita' 08_Materie NON superato (" + str(len(errors)) + " problemi)")
        for e in errors:
            print("  - " + e)
        return 1

    print(f"OK: 08_Materie - TASSONOMIA byte-identica a subjects.csv (live, "
          f"{len(expected_csv)} materie); timing = self-consistency vs snapshot sources.py "
          f"(NON un link Java live, vedi docstring); copertura OK ({real_incarichi} incarichi "
          f"reali, {len(real_subjects)} materie di insegnamento tutte classificate).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
