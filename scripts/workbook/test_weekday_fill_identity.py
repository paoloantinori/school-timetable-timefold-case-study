#!/usr/bin/env python3
"""Identity gate for the weekday-fill policy (golden mining r2 section D CORRECTED,
sheet 25_WeekdayFill). Mirrors the sibling gates. Two tiers:

1. **Sanity invariants:** the derivation returns exactly the 4 corrected-mining cells
   (all G04, Mon-Thu; Friday is 97% and must NOT appear); representative numbers
   match byte for byte; every row respects both the floor (< 80%) and the evidence
   minimum (>= 10 expected days); and the FIRST-PASS artifacts are ABSENT (G12
   Friday, G10/G11/G13/G07 Fridays and G08 Monday: all 83-100% under the
   corrected baseline, they only looked under-filled when stage-day hours counted as
   expected aula+lab).
2. **Round-trip emission:** build, reload, read 25_WeekdayFill BY HEADER NAME, match.

Run: scripts/workbook/.venv/bin/python scripts/workbook/test_weekday_fill_identity.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import weekday_fill_policy as WFP
from generate_workbook import build_workbook

EXPECTED_CELLS = 4  # corrected mining section D

PINNED = [
    ("G04", "Lunedi", 134, 196),
    ("G04", "Martedi", 135, 198),
    ("G04", "Mercoledi", 166, 210),
    ("G04", "Giovedi", 145, 216),
]

# First-pass artifacts (buggy baseline: S-day hours as aula+lab expectation).
MUST_BE_ABSENT = [
    ("G12", "Venerdi"), ("G10", "Venerdi"), ("G11", "Venerdi"),
    ("G13", "Venerdi"), ("G07", "Venerdi"), ("G08", "Lunedi"),
    ("G04", "Venerdi"),
]


def fail(msg: str) -> None:
    print(f"FAIL {msg}")
    sys.exit(1)


def main() -> None:
    rows = WFP.underfill_cells()

    # tier 1: sanity invariants
    if len(rows) != EXPECTED_CELLS:
        fail(f"expected {EXPECTED_CELLS} under-fill cells, got {len(rows)}: "
             f"{[(r['classe'], r['weekday']) for r in rows]}")
    present = {(r["classe"], r["weekday"]) for r in rows}
    for pin in PINNED:
        match = [r for r in rows if (r["classe"], r["weekday"]) == pin[:2]]
        if not match:
            fail(f"pinned cell missing: {pin}")
        if match[0]["oreOsservate"] != pin[2] or match[0]["oreAttese"] != pin[3]:
            fail(f"pinned cell drifted: {pin} -> {match[0]}")
    for r in rows:
        if r["fill"] >= WFP.FILL_FLOOR:
            fail(f"{r['classe']}/{r['weekday']} fill {r['fill']:.2f} not under floor")
        if r["giorniAttesi"] < WFP.MIN_DAYS:
            fail(f"{r['classe']}/{r['weekday']} only {r['giorniAttesi']} expected days")
    for absent in MUST_BE_ABSENT:
        if absent in present:
            fail(f"buggy-baseline artifact present: {absent}")
    print(f"tier 1 OK: {EXPECTED_CELLS} cells, pins match, {len(MUST_BE_ABSENT)}"
          " artifacts absent")

    # tier 2: round-trip emission
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "wb.xlsx"
        wb = build_workbook()
        wb.save(path)
        reloaded = load_workbook(path, read_only=True)
        if "25_WeekdayFill" not in reloaded.sheetnames:
            fail("sheet 25_WeekdayFill missing from the workbook")
        ws = reloaded["25_WeekdayFill"]
        it = ws.iter_rows(values_only=True)
        header = list(next(it))
        col = {name: header.index(name) for name in
               ("classe", "weekday", "oreOsservate", "oreAttese")}
        sheet_rows = []
        for values in it:
            if values[col["classe"]] is None:
                continue
            sheet_rows.append((values[col["classe"]], values[col["weekday"]],
                               values[col["oreOsservate"]], values[col["oreAttese"]]))
        reloaded.close()
    expected_rows = sorted((r["classe"], r["weekday"], r["oreOsservate"], r["oreAttese"])
                           for r in rows)
    if sorted(sheet_rows) != expected_rows:
        fail(f"round-trip drift: sheet {len(sheet_rows)} vs derivation {len(expected_rows)}")
    print(f"tier 2 OK: round-trip emission matches ({len(sheet_rows)} rows)")
    print("PASS weekday-fill identity")


if __name__ == "__main__":
    main()
